import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import os

class ExcelUtils:
    def __init__(self, file_path):
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Excel file not found: {file_path}")
        self.file_path = file_path
        self.workbook = openpyxl.load_workbook(file_path)

    def get_sheet(self, sheet_name=None):
        """Return a sheet by name, or the active sheet if None."""
        if sheet_name:
            if sheet_name in self.workbook.sheetnames:
                return self.workbook[sheet_name]
            else:
                raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        return self.workbook.active

    def create_or_get_sheet(self, sheet_name):
        """Create sheet if not exists, else return existing one."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        return self.workbook.create_sheet(title=sheet_name)

    def read_cell(self, sheet_name, row, col):
        """Read value from a specific cell."""
        sheet = self.get_sheet(sheet_name)
        return sheet.cell(row=row, column=col).value

    def write_cell(self, sheet_name, row, col, value, style=None):
        """Write value into a specific cell with optional style."""
        sheet = self.get_sheet(sheet_name)
        cell = sheet.cell(row=row, column=col)
        cell.value = value

        if style:
            if "font" in style: cell.font = style["font"]
            if "fill" in style: cell.fill = style["fill"]
            if "border" in style: cell.border = style["border"]
            if "alignment" in style: cell.alignment = style["alignment"]

        self.workbook.save(self.file_path)

    def read_all(self, sheet_name):
        """Read all rows and columns from a sheet."""
        sheet = self.get_sheet(sheet_name)
        return [list(r) for r in sheet.iter_rows(values_only=True)]

    def write_matrix(self, sheet_name, matrix, style=None):
        """Write a 2D list (matrix) into a sheet starting at A1."""
        sheet = self.get_sheet(sheet_name)
        for r_idx, row in enumerate(matrix, start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = sheet.cell(row=r_idx, column=c_idx)
                cell.value = value
                if style:
                    if "font" in style: cell.font = style["font"]
                    if "fill" in style: cell.fill = style["fill"]
                    if "border" in style: cell.border = style["border"]
                    if "alignment" in style: cell.alignment = style["alignment"]
        self.workbook.save(self.file_path)

    def clear_sheet(self, sheet_name):
        """Clear all data from a sheet."""
        sheet = self.get_sheet(sheet_name)
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = None
        self.workbook.save(self.file_path)

    def get_total_rows(self, sheet_name):
        """Return total number of rows in a sheet."""
        sheet = self.get_sheet(sheet_name)
        return sheet.max_row

    def get_total_columns(self, sheet_name):
        """Return total number of columns in a sheet."""
        sheet = self.get_sheet(sheet_name)
        return sheet.max_column

    def save(self):
        """Save workbook explicitly."""
        self.workbook.save(self.file_path)
